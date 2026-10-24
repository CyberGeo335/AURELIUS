"""Bounded document intake. Extracted content is evidence, never agent instructions."""
import base64
import csv
import hashlib
import io
from pathlib import PurePath
import zipfile

MAX_BYTES = 20 * 1024 * 1024


def parse_document(filename, encoded):
    try:
        return _parse_document(filename,encoded)
    except (ValueError,TypeError,UnicodeError,zipfile.BadZipFile,KeyError,csv.Error) as exc:
        raise ValueError('Invalid or unsupported document contents') from exc


def _parse_document(filename, encoded):
    if not filename or len(filename)>255 or PurePath(filename).name != filename or '/' in filename or '\\' in filename:
        raise ValueError('A plain filename is required')
    try:
        content=base64.b64decode(encoded,validate=True)
    except (ValueError,TypeError) as exc:
        raise ValueError('Invalid base64 document') from exc
    if not content or len(content)>MAX_BYTES:
        raise ValueError('Document must be 1 byte..20 MB')
    suffix=PurePath(filename).suffix.lower()
    extracted={'status':'REQUIRES_MANUAL_REVIEW','tables':[],'text':None,'is_verified':False}
    if suffix=='.csv':
        try:
            text=content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text=content.decode('cp1251')
        try:
            dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
        except csv.Error:
            dialect=csv.excel
        rows=[]
        for row in csv.reader(io.StringIO(text),dialect):
            if len(rows)>=10000 or len(row)>200 or any(len(cell)>12000 for cell in row):
                raise ValueError('CSV extraction limits exceeded')
            rows.append(row)
        extracted.update(status='EXTRACTED_UNVERIFIED',tables=[{'sheet':'CSV','rows':rows}])
        media='text/csv'
    elif suffix=='.xlsx':
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist())>2000 or sum(item.file_size for item in archive.infolist())>50*1024*1024:
                raise ValueError('Expanded spreadsheet exceeds limits')
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True,keep_links=False)
        try:
            count=0
            for sheet in workbook:
                rows=[]
                if len(extracted['tables'])>=20:
                    raise ValueError('Too many sheets')
                for row in sheet.iter_rows(values_only=True):
                    count+=1
                    if count>10000 or len(row)>200:
                        raise ValueError('Spreadsheet extraction limits exceeded')
                    rows.append([str(v)[:12000] if v is not None else None for v in row])
                extracted['tables'].append({'sheet':sheet.title,'rows':rows})
        finally:
            workbook.close()
        extracted['status']='EXTRACTED_UNVERIFIED'
        media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif suffix=='.pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('Invalid PDF signature')
        # pypdf is optional; unsupported/encrypted/scanned files remain explicit manual evidence.
        try:
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(content))
            if reader.is_encrypted or len(reader.pages)>100:
                extracted['reason']='ENCRYPTED_OR_PAGE_LIMIT'
            else:
                fragments=[]
                total=0
                for page in reader.pages:
                    fragment=page.extract_text() or ''
                    total+=len(fragment)
                    if total>200000:
                        raise ValueError('PDF extraction limit exceeded')
                    fragments.append(fragment)
                extracted.update(text='\n'.join(fragments),status='EXTRACTED_UNVERIFIED' if total else 'REQUIRES_MANUAL_REVIEW')
        except ImportError:
            extracted['reason']='PDF_EXTRACTOR_UNAVAILABLE'
        except Exception as exc:
            raise ValueError('PDF text extraction failed; provide an accessible PDF') from exc
        media='application/pdf'
    elif suffix in {'.png','.jpg','.jpeg','.xls'}:
        extracted['reason']='OCR_OR_LEGACY_FORMAT_REQUIRES_MANUAL_REVIEW'
        media='application/octet-stream'
    else:
        raise ValueError('Unsupported document format')
    return content,media,hashlib.sha256(content).hexdigest(),extracted
